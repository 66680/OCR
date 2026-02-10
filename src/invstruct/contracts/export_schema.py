from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from invstruct.schema_version import EXPORT_SCHEMA_VERSION

REQUIRED_COLUMNS = [
    "file_name",
    "sha256",
    "trace_id",
    "merchant",
    "issue_date",
    "total_amount_gross",
    "currency",
    "status",
    "error_code",
    "error_message",
    "warnings_count",
    "parser_version",
    "retry_key",
]


@dataclass
class ValidationResult:
    ok: bool
    errors: list[str] = field(default_factory=list)
    details: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {"ok": self.ok, "errors": self.errors, "details": self.details}


def get_export_contract_schema() -> dict[str, Any]:
    return {
        "$id": f"invstruct/export-schema/v{EXPORT_SCHEMA_VERSION}",
        "schema_version": EXPORT_SCHEMA_VERSION,
        "required_columns": REQUIRED_COLUMNS,
        "notes": [
            "CSV may include comment headers.",
            "XLSX must contain _meta sheet with schema version marker.",
        ],
    }


def _read_csv_columns(path: Path) -> tuple[list[str], list[str], list[str]]:
    if not path.exists():
        return [], [], [f"csv_not_found:{path}"]
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    comment_lines: list[str] = []
    header_line = ""
    for line in lines:
        if line.startswith("#"):
            comment_lines.append(line)
            continue
        if line.strip() == "":
            continue
        header_line = line
        break
    if not header_line:
        return comment_lines, [], ["csv_header_missing"]
    columns = next(csv.reader([header_line]))
    return comment_lines, columns, []


def validate_export_csv(
    path: str | Path,
    expect_version: int = EXPORT_SCHEMA_VERSION,
    header_comments: bool = True,
) -> ValidationResult:
    csv_path = Path(path)
    comment_lines, columns, parse_errors = _read_csv_columns(csv_path)
    errors = list(parse_errors)

    if header_comments:
        expect = f"# invstruct_export_schema_version={expect_version}"
        if expect not in comment_lines:
            errors.append("csv_schema_version_comment_missing")

    missing_required = [column for column in REQUIRED_COLUMNS if column not in columns]
    if missing_required:
        errors.append(f"csv_missing_columns:{','.join(missing_required)}")

    return ValidationResult(
        ok=len(errors) == 0,
        errors=errors,
        details={"columns": columns, "comment_lines": comment_lines},
    )


def validate_export_xlsx(
    path: str | Path,
    expect_version: int = EXPORT_SCHEMA_VERSION,
) -> ValidationResult:
    xlsx_path = Path(path)
    errors: list[str] = []
    if not xlsx_path.exists():
        return ValidationResult(ok=False, errors=[f"xlsx_not_found:{xlsx_path}"])

    try:
        from openpyxl import load_workbook
    except Exception:
        return ValidationResult(ok=False, errors=["xlsx_dependency_missing_openpyxl"])

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        sheet_names = workbook.sheetnames
        if "_meta" not in sheet_names:
            errors.append("xlsx_meta_sheet_missing")
        else:
            meta_sheet = workbook["_meta"]
            meta_key = str(meta_sheet["A1"].value or "")
            meta_value = meta_sheet["B1"].value
            if meta_key != "invstruct_export_schema_version":
                errors.append("xlsx_meta_key_invalid")
            if str(meta_value) != str(expect_version):
                errors.append("xlsx_schema_version_invalid")

        data_sheets = [name for name in sheet_names if name != "_meta"]
        if not data_sheets:
            errors.append("xlsx_data_sheet_missing")
            columns: list[str] = []
        else:
            row = next(workbook[data_sheets[0]].iter_rows(min_row=1, max_row=1, values_only=True))
            columns = [str(item) for item in row if item is not None]
            missing_required = [column for column in REQUIRED_COLUMNS if column not in columns]
            if missing_required:
                errors.append(f"xlsx_missing_columns:{','.join(missing_required)}")
    finally:
        workbook.close()

    return ValidationResult(
        ok=len(errors) == 0,
        errors=errors,
        details={"sheet_names": workbook.sheetnames, "columns": columns if "columns" in locals() else []},
    )
